# The bulk tests' workbook (packages/core/test/table.test.ts, bulk.test.ts), written by openpyxl:
# python3 make_products_xlsx.py products.xlsx
import datetime, sys
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
wb = Workbook()
ws = wb.active
ws.title = "Produtos"
ws.append(["GTIN", "Lote", "Validade", "Peso (kg)", "Descrição", "Fabricação"])
rows = [
    ("09506000134376", "B42", datetime.date(2027, 12, 31), 0.5, "Chocolate & cacau <100 g>"),
    ("09506000134352", "L-7", datetime.date(2028, 1, 5), 1.25, None),
    ("09506000134353", "X1", datetime.date(2027, 6, 30), 2, "bad check digit"),
    ("", "Y", datetime.date(2027, 6, 30), 2, "no GTIN"),
]
for r in rows: ws.append(list(r) + [datetime.date(2026, 10, 1)])
for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
    for c in row: c.number_format = "dd/mm/yyyy"
for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
    for c in row: c.number_format = "mm-dd-yy"  # Excel's built-in date format 14
ws["E3"] = CellRichText([TextBlock(InlineFont(b=True), "Rich"), " text"])
ws["A1"].number_format = "@"
other = wb.create_sheet("Notes")
other.append(["ignored"])
wb.save(sys.argv[1])
